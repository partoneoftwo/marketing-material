"""Build product sheet (PDF), sales deck (PPTX), social images (PNG) and texts.md per language from content/<lang>.json."""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from build_site import LANGS, bowtie_svg  # noqa: E402

from fpdf import FPDF  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

NAVY, TEAL, CORAL, SUN, INK, MUTED = (31, 42, 68), (20, 184, 166), (249, 115, 98), (251, 191, 36), (38, 49, 77), (91, 101, 133)
FONTS = ROOT / "tools" / "fonts"
SCREENS = ROOT / "assets" / "screens"
BOWTIE_PNG = ROOT / "assets" / "bowtie.png"


def rgb(t):
    return RGBColor(*t)


# --- PDF product sheet (A4 portrait) ---------------------------------------------------------------------

def product_sheet(c, out):
    pdf = FPDF(format="A4")
    pdf.add_font("DejaVu", "", str(FONTS / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(FONTS / "DejaVuSans-Bold.ttf"))
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.set_title("SpaceRisk")
    pdf.set_fill_color(*NAVY)
    pdf.rect(0, 0, 210, 58, style="F")
    pdf.image(str(ROOT / "assets" / "icon.png"), 14, 10, 12, 12)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("DejaVu", "B", 15)
    pdf.set_xy(29, 12)
    pdf.cell(60, 8, "SpaceRisk")
    pdf.set_xy(14, 26)
    pdf.set_font("DejaVu", "B", 18)
    pdf.multi_cell(182, 8, c["sheet"]["headline"], align="L")
    pdf.set_x(14)
    pdf.set_font("DejaVu", "", 10)
    pdf.set_text_color(214, 222, 245)
    pdf.multi_cell(182, 5, c["sheet"]["sub"], align="L")
    pdf.image(str(BOWTIE_PNG), 14, 62, 182)
    y = 62 + 182 * 0.42 + 4
    pdf.set_text_color(*NAVY)
    pdf.set_font("DejaVu", "B", 12)
    pdf.set_xy(14, y)
    pdf.cell(0, 6, c["sheet"]["flow_title"])
    y += 8
    for lane, steps, colour in ((c["flow"]["lane_team"], c["flow"]["steps_team"], TEAL), (c["flow"]["lane_project"], c["flow"]["steps_project"], CORAL)):
        pdf.set_fill_color(*colour)
        pdf.rect(14, y, 2, 16, style="F")
        x = 19
        for i, s in enumerate(steps):
            pdf.set_fill_color(246, 248, 252)
            pdf.rect(x, y, 54, 16, style="F")
            pdf.set_xy(x + 2, y + 1.5)
            pdf.set_font("DejaVu", "B", 8)
            pdf.set_text_color(*NAVY)
            pdf.cell(50, 4, s["who"])
            pdf.set_xy(x + 2, y + 6)
            pdf.set_font("DejaVu", "", 6.5)
            pdf.set_text_color(*MUTED)
            pdf.multi_cell(50, 3, s["what"], align="L")
            if i < len(steps) - 1:
                pdf.set_fill_color(*colour)
                pdf.ellipse(x + 55, y + 7, 3, 3, style="F")
            x += 59
        y += 19
    pdf.set_xy(14, y + 1)
    pdf.set_font("DejaVu", "B", 12)
    pdf.set_text_color(*NAVY)
    pdf.cell(0, 6, c["sheet"]["features_title"])
    y += 9
    accents = [TEAL, CORAL, SUN, (167, 139, 250)]
    for i, f in enumerate(c["features"]["items"]):
        col, row = i % 2, i // 2
        x, yy = 14 + col * 92, y + row * 15
        pdf.set_fill_color(*accents[i % 4])
        pdf.rect(x, yy, 1.6, 12, style="F")
        pdf.set_xy(x + 3.5, yy)
        pdf.set_font("DejaVu", "B", 8.5)
        pdf.set_text_color(*NAVY)
        pdf.cell(86, 4, f["title"])
        pdf.set_xy(x + 3.5, yy + 4.5)
        pdf.set_font("DejaVu", "", 6.8)
        pdf.set_text_color(*MUTED)
        pdf.multi_cell(86, 3.1, f["text"], align="L")
    y += 4 * 15
    room = 281 - y  # fill what is left above the footer with three small screens
    if room > 14:
        for i, name in enumerate(("domain-bowtie.png", "executive-dashboard.png", "financial-perspective.png")):
            pdf.image(str(SCREENS / name), 14 + i * 61.5, y, w=59, h=min(room, 59 * 0.6), keep_aspect_ratio=True)
    pdf.set_fill_color(*NAVY)
    pdf.rect(0, 284, 210, 13, style="F")
    pdf.set_xy(14, 287)
    pdf.set_font("DejaVu", "B", 9)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(182, 6, c["sheet"]["cta"], align="C")
    out.write_bytes(bytes(pdf.output()))


# --- PPTX sales deck (16:9) ------------------------------------------------------------------------------

def sales_deck(c, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, value, size, bold=False, colour=INK):
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        for i, line in enumerate(value if isinstance(value, list) else [value]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text, p.font.size, p.font.bold, p.font.color.rgb = line, Pt(size), bold, rgb(colour)
            p.space_after = Pt(10)

    def band(slide, colour):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.22))
        s.fill.solid(); s.fill.fore_color.rgb = rgb(colour); s.line.fill.background()

    # cover
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb(NAVY); bg.line.fill.background()
    s.shapes.add_picture(str(ROOT / "assets" / "icon.png"), Inches(0.7), Inches(0.6), Inches(0.6))
    text(s, 1.45, 0.62, 5, 0.6, "SpaceRisk", 24, True, (255, 255, 255))
    text(s, 0.7, 1.7, 6.2, 2.6, c["hero"]["title"], 36, True, (255, 255, 255))
    text(s, 0.7, 4.5, 6.0, 2, c["sheet"]["sub"], 16, False, (199, 208, 239))
    s.shapes.add_picture(str(BOWTIE_PNG), Inches(7.1), Inches(1.9), Inches(5.8))
    accents = [CORAL, TEAL, SUN, (167, 139, 250)]
    for i, sl in enumerate(c["deck"]["slides"]):
        s = prs.slides.add_slide(blank)
        band(s, accents[i % 4])
        text(s, 0.7, 0.55, 12, 1.1, sl["title"], 32, True, NAVY)
        image = sl.get("image")
        if i == 2:  # bowtie slide: the drawn bowtie, plus a real screen
            s.shapes.add_picture(str(BOWTIE_PNG), Inches(0.7), Inches(1.9), Inches(6.2))
            text(s, 0.7, 5.0, 6.2, 2, sl["bullets"], 15, False, MUTED)
            s.shapes.add_picture(str(SCREENS / image), Inches(7.3), Inches(1.9), Inches(5.4))
        elif image:
            text(s, 0.7, 1.9, 5.2, 4.8, ["• " + b for b in sl["bullets"]], 19, False, INK)
            s.shapes.add_picture(str(SCREENS / image), Inches(6.2), Inches(1.8), Inches(6.5))
        elif i in (3,):  # flow slide: steps as connected cards
            steps = c["flow"]["steps_team"]
            for k, st in enumerate(steps):
                card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7 + k * 4.2), Inches(2.4), Inches(3.6), Inches(2.4))
                card.fill.solid(); card.fill.fore_color.rgb = rgb((227, 251, 246)); card.line.color.rgb = rgb(TEAL)
                text(s, 0.95 + k * 4.2, 2.6, 3.2, 0.6, st["who"], 22, True, NAVY)
                text(s, 0.95 + k * 4.2, 3.3, 3.2, 1.4, st["what"], 15, False, MUTED)
                if k < 2:
                    arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.35 + k * 4.2), Inches(3.35), Inches(0.5), Inches(0.5))
                    arr.fill.solid(); arr.fill.fore_color.rgb = rgb(TEAL); arr.line.fill.background()
            text(s, 0.7, 5.3, 12, 1, c["flow"]["barrier_note"], 16, False, MUTED)
        else:
            single = len(sl["bullets"]) == 1
            for k, b in enumerate(sl["bullets"]):
                card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.9 + k * 1.15), Inches(11.9),
                                          Inches(2.6 if single else 0.95))
                card.fill.solid(); card.fill.fore_color.rgb = rgb((246, 248, 252)); card.line.color.rgb = rgb(accents[(i + k) % 4])
                text(s, 1.0, 2.25 if single else 2.05 + k * 1.15, 11.4, 2.2 if single else 0.8, b, 28 if single else 20, single, INK)
    prs.save(out)


# --- social images (rendered HTML) ------------------------------------------------------------------------

SOCIAL = """<!doctype html><html><head><meta charset="utf-8"><link rel="stylesheet" href="{css}"><style>
body{{margin:0;width:{w}px;height:{h}px;overflow:hidden;background:linear-gradient(135deg,#1f2a44 0%,#26406b 60%,#14b8a6 140%);color:#fff;
font-family:system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;display:flex;flex-direction:{dir};align-items:center;gap:{gap}px;padding:{pad}px;box-sizing:border-box}}
.t{{flex:1}} .b{{display:flex;align-items:center;gap:12px;font-weight:800;font-size:{bs}px}} .b img{{width:{bs}px;height:{bs}px;border-radius:8px}}
h1{{font-size:{hs}px;line-height:1.08;margin:{hm}px 0 14px;color:#fff;text-align:left}} p{{font-size:{ps}px;color:#cbd5f5;margin:0}}
.art{{flex:{af};background:#fff;border-radius:26px;padding:14px;width:{aw}}}</style></head><body>
<div class="t"><div class="b"><img src="{icon}">SpaceRisk</div><h1>{headline}</h1><p>{sub}</p></div><div class="art">{svg}</div></body></html>"""


def social_images(c, outdir, page):
    outdir.mkdir(parents=True, exist_ok=True)
    for name, w, h, d, pad, hs, ps, bs, af, aw in (("linkedin-1200x627.png", 1200, 627, "row", 56, 50, 22, 28, "0 0 520px", "520px"),
                                                   ("x-1500x500.png", 1500, 500, "row", 50, 46, 22, 26, "0 0 620px", "620px"),
                                                   ("square-1080.png", 1080, 1080, "column", 70, 64, 28, 34, "0 0 auto", "100%")):
        html = SOCIAL.format(css=(ROOT / "assets" / "style.css").as_uri(), icon=(ROOT / "assets" / "icon.png").as_uri(), w=w, h=h, dir=d,
                             gap=40, pad=pad, hs=hs, hm=18, ps=ps, bs=bs, af=af, aw=aw, headline=c["social"]["headline"],
                             sub=c["social"]["sub"], svg=bowtie_svg("social"))
        tmp = outdir / ".social.html"
        tmp.write_text(html)
        page.set_viewport_size({"width": w, "height": h})
        page.goto(tmp.as_uri())
        page.wait_for_timeout(900)
        page.screenshot(path=str(outdir / name))
        tmp.unlink()


def texts_md(c, out):
    t = c["texts"]
    lines = [f"# SpaceRisk · {c['lang_name']}", "", "## Taglines", *[f"- {x}" for x in t["taglines"]], "",
             "## Short description", t["short"], "", "## Medium description", t["medium"], "", "## Long description", t["long"], "",
             "## Feature blurbs", *[f"- **{f['title']}:** {f['text']}" for f in c["features"]["items"]], "", "## Boilerplate", t["boilerplate"], ""]
    out.write_text("\n".join(lines))


def main():
    sys.path.insert(0, "/tmp/pwlib")
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(device_scale_factor=2)
        # the bowtie illustration used in the PDF and deck, rendered once from the site's SVG
        page.set_viewport_size({"width": 1000, "height": 420})
        css = (ROOT / "assets" / "style.css").read_text()
        page.set_content(f'<html><head><style>{css} .bt .pt {{ display: none !important; }}</style></head>'
                         f'<body style="margin:0;background:#fff">{bowtie_svg("png")}</body></html>')
        page.wait_for_timeout(1200)
        page.locator("svg").screenshot(path=str(BOWTIE_PNG))
        for lang in LANGS:
            c = json.loads((ROOT / "content" / f"{lang}.json").read_text())
            out = ROOT / "materials" / lang
            out.mkdir(parents=True, exist_ok=True)
            product_sheet(c, out / "spacerisk-product-sheet.pdf")
            sales_deck(c, out / "spacerisk-sales-deck.pptx")
            texts_md(c, out / "texts.md")
            social_images(c, out / "social", page)
            print("built", lang)
        browser.close()


if __name__ == "__main__":
    main()
